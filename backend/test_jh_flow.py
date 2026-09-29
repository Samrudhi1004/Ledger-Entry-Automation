"""
test_jh_flow.py

End-to-End Automated Test Suite for Autonomous Maintenance (Jishu Hozen - JH)
Inspection Lifecycle:
1. Checklist Template Generation & Download (Form QF/MF-08)
2. Checklist File Upload & Parsing (Excel & Unicode/Hindi text)
3. Checklist Bulk Save & Version Control (v1 -> v2 -> rollback)
4. Checklist Active Items Retrieval for Operators
5. Shift Inspection Submission (ALL_OK, HAS_ISSUES, CORRECTED, and Idempotency)
6. Inspection History & Detailed Reports Filtering
7. Monthly 31-Day Monitoring Matrix API (3-shift grid)
8. Final Official Output Generation (Form QF/MF-08 Excel Export)
"""

import os
import io
import sys
import django
from datetime import date

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

import openpyxl
from rest_framework.test import APIRequestFactory, force_authenticate
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.users.models import User
from apps.machines.models import Machine, Plant, Factory
from apps.inspections.models import (
    JHChecklistVersion,
    JHChecklistItem,
    JHInspectionRecord,
    JHInspectionItemResult,
)
from apps.inspections.jh_checklist_parser import (
    generate_jh_template_xlsx,
    parse_jh_excel,
    parse_jh_checklist_file,
)
from apps.inspections.jh_excel_generator import generate_jh_matrix_xlsx
from apps.inspections.views import (
    JHChecklistItemsView,
    JHChecklistUploadParseView,
    JHChecklistBulkSaveView,
    JHChecklistTemplateDownloadView,
    JHChecklistVersionListView,
    JHChecklistVersionDetailView,
    JHChecklistVersionRestoreView,
    JHChecklistStatusView,
    JHInspectionSubmitView,
    JHInspectionReportsView,
    JHInspectionDetailView,
    JHInspectionMatrixView,
    JHInspectionMatrixExportExcelView,
)

TAG = "[TEST_JH_FLOW] "


class JHTestRunner:
    def __init__(self):
        self.factory = APIRequestFactory()
        self.passed = 0
        self.failed = 0
        self.tests = []

    def assert_test(self, condition: bool, test_id: str, description: str, detail: str = ""):
        if condition:
            self.passed += 1
            print(f"  [PASS] {test_id}: {description}")
            self.tests.append({"id": test_id, "desc": description, "status": "PASS", "detail": detail})
        else:
            self.failed += 1
            print(f"  [FAIL] {test_id}: {description} --> {detail}")
            self.tests.append({"id": test_id, "desc": description, "status": "FAIL", "detail": detail})

    def run_all(self):
        print("\n" + "=" * 80)
        print("  JISHU HOZEN (JH) INSPECTION: COMPLETE END-TO-END FLOW TEST SUITE")
        print("=" * 80 + "\n")

        self.setup_fixtures()

        try:
            self.test_phase1_template_download()
            self.test_phase2_file_parsing()
            self.test_phase3_upload_parse_api()
            self.test_phase4_bulk_save_and_versioning()
            self.test_phase5_checklist_items_retrieval()
            self.test_phase6_shift_inspection_submissions()
            self.test_phase7_reports_and_history()
            self.test_phase8_matrix_grid_api()
            self.test_phase9_final_excel_export()
            self.test_phase10_multimachine_isolation()
        finally:
            self.cleanup()

        print("\n" + "-" * 80)
        print(f"  TEST SUMMARY: {self.passed} Passed, {self.failed} Failed (Total: {self.passed + self.failed})")
        print("-" * 80 + "\n")

        return self.failed == 0

    def setup_fixtures(self):
        print(f"{TAG}Setting up clean isolated fixtures...")
        # 1. Organization & Factory
        self.factory_obj, _ = Factory.objects.get_or_create(
            code="JH-FAC-TEST",
            defaults={"name": "Liha Tech Precision Plant", "is_active": True}
        )
        self.plant_obj, _ = Plant.objects.get_or_create(
            code="JH-PLT-TEST",
            factory=self.factory_obj,
            defaults={"name": "Machining Shop Floor"}
        )
        self.machine_obj, _ = Machine.objects.get_or_create(
            machine_code="JH-MCH-01",
            defaults={"name": "HASS VMC-01", "plant": self.plant_obj, "status": "active"}
        )

        # 2. Users
        self.admin_user, _ = User.objects.get_or_create(
            username="jh_test_admin",
            defaults={
                "role": "admin",
                "is_superuser": True,
                "first_name": "Admin",
                "last_name": "User",
                "access_grants": ["production.jh.manage", "production.jh.view", "production.jh.submit"]
            }
        )
        self.admin_user.is_superuser = True
        self.admin_user.access_grants = ["production.jh.manage", "production.jh.view", "production.jh.submit"]
        self.admin_user.set_password("pass123")
        self.admin_user.save()

        self.operator_user, _ = User.objects.get_or_create(
            username="jh_test_operator",
            defaults={
                "role": "operator",
                "first_name": "Suresh",
                "last_name": "Kadam",
                "assigned_shift": "ALL",
                "access_grants": ["production.jh.submit", "production.jh.view"]
            }
        )
        self.operator_user.assigned_shift = "ALL"
        self.operator_user.access_grants = ["production.jh.submit", "production.jh.view"]
        self.operator_user.set_password("pass123")
        self.operator_user.save()

        # Clean any prior test JH records and checklists
        self.cleanup_all_test_data()

    def cleanup_all_test_data(self):
        test_machines = Machine.objects.filter(machine_code__in=["JH-MCH-01", "JH-MCH-02", "JH-MCH-03"])
        JHInspectionRecord.objects.filter(machine__in=test_machines).delete()
        JHChecklistItem.objects.filter(machine__in=test_machines).delete()
        JHChecklistItem.objects.filter(sub_no__startswith="TEST-").delete()
        JHChecklistItem.objects.filter(sub_no__startswith="MCH").delete()
        JHChecklistVersion.objects.filter(machine__in=test_machines).delete()
        JHChecklistVersion.objects.filter(filename__contains="test_").delete()
        JHChecklistVersion.objects.filter(filename__contains="cnc_").delete()

    def cleanup(self):
        print(f"{TAG}Cleaning up test records...")
        self.cleanup_all_test_data()

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 1: TEMPLATE DOWNLOAD
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase1_template_download(self):
        print("\n--- PHASE 1: Checklist Template Generation & Download ---")

        # 1.1 Direct Generator Function
        buf = generate_jh_template_xlsx()
        wb = openpyxl.load_workbook(filename=buf)
        ws = wb.active
        self.assert_test(
            ws.title == "JH_Checklist_Template",
            "TC-JH-01A",
            "Template workbook has 'JH_Checklist_Template' sheet title"
        )
        self.assert_test(
            "FORM QF/MF-08" in str(ws["A1"].value),
            "TC-JH-01B",
            "Template banner contains Form QF/MF-08 title"
        )
        self.assert_test(
            ws.max_row >= 8,
            "TC-JH-01C",
            "Template contains headers and sample pre-filled checkpoints",
            f"Row count: {ws.max_row}"
        )

        # 1.2 HTTP API View
        view = JHChecklistTemplateDownloadView.as_view()
        request = self.factory.get('/api/inspections/jh/checklist/template/')
        force_authenticate(request, user=self.admin_user)
        response = view(request)

        self.assert_test(
            response.status_code == 200,
            "TC-JH-01D",
            "GET /api/inspections/jh/checklist/template/ returns HTTP 200"
        )
        self.assert_test(
            "spreadsheetml.sheet" in response.get('Content-Type', ''),
            "TC-JH-01E",
            "Template download returns proper Excel MIME type"
        )
        self.assert_test(
            "JH_Checklist_Template" in response.get('Content-Disposition', ''),
            "TC-JH-01F",
            "Template download provides filename attachment",
            f"Header: {response.get('Content-Disposition', '')}"
        )

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 2: PARSER CORE (Excel & Hindi Unicode support)
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase2_file_parsing(self):
        print("\n--- PHASE 2: File Parsing & Data Normalization ---")

        # Create sample Excel workbook in memory with English & Hindi
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Sub No", "Assembly", "Sub Assembly", "Check Point", "Standard", "Tool", "Rank", "Freq", "Timing", "Action"])
        ws.append(["TEST-1.1", "1. Machine Front Side", "एफ एम एफ बोर्ड", "पूरा एफ एम एफ बोर्ड साफ करो", "धूल और तेल से मुक्त", "VISUAL", "B", "D", "5 DPT", "Clean, Inspect"])
        ws.append(["TEST-1.2", "1. Machine Front Side", "एफ आर एल", "एफ आर एल का एयर प्रेशर चेक करो", "5-6 bar", "VISUAL", "D", "D", "5 DPT", "Inspect"])
        ws.append(["TEST-2.1", "2. FIXTURE", "फिक्सचर", "गाइड पिन ऑयलिंग करो", "ऑयलिंग ठीक हो", "TOOL", "D", "D", "10 DPT", "Lubricate"])
        ws.append(["TEST-2.2", "2. FIXTURE", "फिक्सचर", "माउंटिंग बोल्ट टाइट करो", "टाइट होना चाहिए", "TOOL", "A", "D", "15 DPT", "Retighten"])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        file_bytes = buf.getvalue()

        items = parse_jh_excel(file_bytes)
        self.assert_test(
            len(items) == 4,
            "TC-JH-02A",
            "parse_jh_excel accurately parses all 4 rows",
            f"Parsed: {len(items)}"
        )

        item1 = items[0]
        self.assert_test(
            item1["sub_no"] == "TEST-1.1" and "एफ एम एफ" in item1["check_point"],
            "TC-JH-02B",
            "Parser preserves pure Hindi Unicode checkpoints without corruption"
        )
        self.assert_test(
            item1["tool_type"] == "VISUAL" and item1["action_clean"] is True,
            "TC-JH-02C",
            "Parser detects tool type (VISUAL) and action (Clean)"
        )

        item3 = items[2]
        self.assert_test(
            item3["tool_type"] == "TOOL" and item3["action_lubricate"] is True,
            "TC-JH-02D",
            "Parser detects tool type (TOOL) and action (Lubricate)"
        )

        item4 = items[3]
        self.assert_test(
            item4["action_retighten"] is True and item4["rank"] == "A",
            "TC-JH-02E",
            "Parser detects Retighten action and Rank A"
        )

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 3: UPLOAD & PARSE API
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase3_upload_parse_api(self):
        print("\n--- PHASE 3: Upload & Parse API Endpoint ---")

        # 3.1 Valid upload
        template_buf = generate_jh_template_xlsx()
        uploaded_file = SimpleUploadedFile(
            "test_jh_checklist.xlsx",
            template_buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        view = JHChecklistUploadParseView.as_view()
        request = self.factory.post('/api/inspections/jh/checklist/upload_parse/', {'file': uploaded_file}, format='multipart')
        force_authenticate(request, user=self.admin_user)
        response = view(request)

        self.assert_test(
            response.status_code == 200,
            "TC-JH-03A",
            "POST /api/inspections/jh/checklist/upload_parse/ returns HTTP 200"
        )
        self.assert_test(
            response.data.get("success") is True and response.data.get("count", 0) >= 5,
            "TC-JH-03B",
            "Upload response returns success flag and parsed items count",
            f"Items parsed: {response.data.get('count')}"
        )

        # 3.2 Missing file error
        request_empty = self.factory.post('/api/inspections/jh/checklist/upload_parse/', {}, format='multipart')
        force_authenticate(request_empty, user=self.admin_user)
        res_empty = view(request_empty)
        self.assert_test(
            res_empty.status_code == 400,
            "TC-JH-03C",
            "POST upload_parse without file returns HTTP 400 Bad Request"
        )

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 4: BULK SAVE & VERSION CONTROL
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase4_bulk_save_and_versioning(self):
        print("\n--- PHASE 4: Checklist Bulk Save & Version Control ---")

        # 4.1 Save Version 1 with 3 items
        items_v1 = [
            {
                "sub_no": "TEST-1.1",
                "assembly": "1. Front Side",
                "sub_assembly": "Board",
                "check_point": "Clean board",
                "standard": "Clean and dry",
                "tool_type": "VISUAL",
                "rank": "D",
                "frequency": "D",
                "timing_sec": "5 DPT",
                "action_clean": True,
                "action_inspect": True,
                "sort_order": 1
            },
            {
                "sub_no": "TEST-1.2",
                "assembly": "1. Front Side",
                "sub_assembly": "Air Gauge",
                "check_point": "Check pressure",
                "standard": "5-6 bar",
                "tool_type": "VISUAL",
                "rank": "B",
                "frequency": "D",
                "timing_sec": "5 DPT",
                "action_inspect": True,
                "sort_order": 2
            },
            {
                "sub_no": "TEST-2.1",
                "assembly": "2. Fixture",
                "sub_assembly": "Pins",
                "check_point": "Oil fixture pins",
                "standard": "Lubricated",
                "tool_type": "TOOL",
                "rank": "D",
                "frequency": "D",
                "timing_sec": "10 DPT",
                "action_lubricate": True,
                "sort_order": 3
            },
        ]

        view_save = JHChecklistBulkSaveView.as_view()
        req_v1 = self.factory.post(
            '/api/inspections/jh/checklist/bulk_save/',
            {"filename": "test_v1.xlsx", "items": items_v1, "notes": "Initial setup"},
            format='json'
        )
        force_authenticate(req_v1, user=self.admin_user)
        res_v1 = view_save(req_v1)

        self.assert_test(
            res_v1.status_code == 200,
            "TC-JH-04A",
            "Bulk save creates Version 1 returning HTTP 200 OK"
        )
        v1_num = res_v1.data.get("version_number")
        self.assert_test(
            v1_num is not None,
            "TC-JH-04B",
            f"Version 1 assigned number {v1_num}"
        )

        active_count_v1 = JHChecklistItem.objects.filter(is_active=True, sub_no__startswith="TEST-").count()
        self.assert_test(
            active_count_v1 == 3,
            "TC-JH-04C",
            "Version 1 has 3 active items in database",
            f"Count: {active_count_v1}"
        )

        # 4.2 Save Version 2 with 2 items (verifying previous version is deactivated)
        items_v2 = [
            {
                "sub_no": "TEST-1.1",
                "assembly": "1. Front Side",
                "check_point": "Clean board updated",
                "standard": "Spotless",
                "tool_type": "VISUAL",
                "rank": "A",
                "sort_order": 1
            },
            {
                "sub_no": "TEST-3.1",
                "assembly": "3. Hydraulic",
                "check_point": "Check pump temperature",
                "standard": "< 50 deg C",
                "tool_type": "TOUCH",
                "rank": "D",
                "sort_order": 2
            }
        ]

        req_v2 = self.factory.post(
            '/api/inspections/jh/checklist/bulk_save/',
            {"filename": "test_v2.xlsx", "items": items_v2, "notes": "Revised checkpoints"},
            format='json'
        )
        force_authenticate(req_v2, user=self.admin_user)
        res_v2 = view_save(req_v2)
        v2_num = res_v2.data.get("version_number")

        self.assert_test(
            v2_num == v1_num + 1,
            "TC-JH-04D",
            "Version number increments monotonically (v2 = v1 + 1)"
        )

        active_count_v2 = JHChecklistItem.objects.filter(is_active=True, sub_no__startswith="TEST-").count()
        self.assert_test(
            active_count_v2 == 2,
            "TC-JH-04E",
            "Version 2 replaces active items: old items deactivated, 2 new items active",
            f"Active count: {active_count_v2}"
        )

        # 4.3 Version History & Rollback/Restore
        view_restore = JHChecklistVersionRestoreView.as_view()
        req_restore = self.factory.post(f'/api/inspections/jh/checklist/versions/{v1_num}/restore/')
        force_authenticate(req_restore, user=self.admin_user)
        res_restore = view_restore(req_restore, version_number=v1_num)

        self.assert_test(
            res_restore.status_code == 200,
            "TC-JH-04F",
            "Version restore endpoint returns HTTP 200"
        )
        active_count_restored = JHChecklistItem.objects.filter(is_active=True, sub_no__startswith="TEST-").count()
        self.assert_test(
            active_count_restored == 3,
            "TC-JH-04G",
            "Rollback restored Version 1: exactly 3 items reactivated",
            f"Active count: {active_count_restored}"
        )

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 5: CHECKLIST ITEMS RETRIEVAL (Operator View)
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase5_checklist_items_retrieval(self):
        print("\n--- PHASE 5: Checklist Items Retrieval for Mobile & Web ---")

        view = JHChecklistItemsView.as_view()
        request = self.factory.get('/api/inspections/jh/items/')
        force_authenticate(request, user=self.operator_user)
        response = view(request)

        self.assert_test(
            response.status_code == 200,
            "TC-JH-05A",
            "GET /api/inspections/jh/items/ returns HTTP 200 for operator"
        )
        results = response.data.get("results", [])
        test_items = [it for it in results if it.get("sub_no", "").startswith("TEST-")]
        self.assert_test(
            len(test_items) == 3,
            "TC-JH-05B",
            "Returns exactly active checkpoints",
            f"Items found: {len(test_items)}"
        )
        self.assert_test(
            test_items[0]["sort_order"] <= test_items[1]["sort_order"],
            "TC-JH-05C",
            "Checkpoints are sorted by sort_order"
        )

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 6: SHIFT INSPECTION SUBMISSIONS (Execution Lifecycle)
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase6_shift_inspection_submissions(self):
        print("\n--- PHASE 6: Shift Inspection Submission & Status Calculations ---")

        active_items = list(JHChecklistItem.objects.filter(is_active=True, sub_no__startswith="TEST-").order_by('sort_order'))
        item1, item2, item3 = active_items[0], active_items[1], active_items[2]

        view_submit = JHInspectionSubmitView.as_view()
        today_str = date.today().isoformat()

        # 6.1 Shift I Submission: ALL_OK
        payload_all_ok = {
            "machine_id": self.machine_obj.id,
            "date": today_str,
            "shift": "I",
            "overall_remarks": "Everything running smoothly",
            "results": [
                {"item_id": item1.id, "status": "OK"},
                {"item_id": item2.id, "status": "OK"},
                {"item_id": item3.id, "status": "OK"},
            ]
        }
        req1 = self.factory.post('/api/inspections/jh/submit/', payload_all_ok, format='json')
        force_authenticate(req1, user=self.operator_user)
        res1 = view_submit(req1)

        self.assert_test(
            res1.status_code == 200 or res1.status_code == 201,
            "TC-JH-06A",
            f"Shift I submission returns success status ({res1.status_code})"
        )
        rec1 = JHInspectionRecord.objects.get(machine=self.machine_obj, date=date.today(), shift="I")
        self.assert_test(
            rec1.status == JHInspectionRecord.Status.ALL_OK and rec1.ok_items == 3 and rec1.not_ok_items == 0,
            "TC-JH-06B",
            "Shift I status correctly calculated as ALL_OK (3 OK, 0 NOT_OK)"
        )

        # 6.2 Shift II Submission: HAS_ISSUES (1 NOT_OK)
        payload_has_issues = {
            "machine_id": self.machine_obj.id,
            "date": today_str,
            "shift": "II",
            "overall_remarks": "Air gauge pressure dropped below 4 bar",
            "results": [
                {"item_id": item1.id, "status": "OK"},
                {"item_id": item2.id, "status": "NOT_OK", "remark": "Low air pressure: 3.5 bar"},
                {"item_id": item3.id, "status": "OK"},
            ]
        }
        req2 = self.factory.post('/api/inspections/jh/submit/', payload_has_issues, format='json')
        force_authenticate(req2, user=self.operator_user)
        res2 = view_submit(req2)

        self.assert_test(
            res2.status_code == 200 or res2.status_code == 201,
            "TC-JH-06C",
            f"Shift II submission returns success status ({res2.status_code})"
        )
        rec2 = JHInspectionRecord.objects.get(machine=self.machine_obj, date=date.today(), shift="II")
        self.assert_test(
            rec2.status == JHInspectionRecord.Status.HAS_ISSUES and rec2.not_ok_items == 1,
            "TC-JH-06D",
            "Shift II status correctly calculated as HAS_ISSUES (1 NOT_OK)"
        )

        # 6.3 Shift III Submission: CORRECTED (Not OK Correction Done)
        payload_corrected = {
            "machine_id": self.machine_obj.id,
            "date": today_str,
            "shift": "III",
            "overall_remarks": "Compressor regulator valve adjusted",
            "results": [
                {"item_id": item1.id, "status": "OK"},
                {"item_id": item2.id, "status": "CORRECTED", "remark": "Pressure set back to 5.5 bar", "action_taken": "Replaced valve seal"},
                {"item_id": item3.id, "status": "OK"},
            ]
        }
        req3 = self.factory.post('/api/inspections/jh/submit/', payload_corrected, format='json')
        force_authenticate(req3, user=self.operator_user)
        res3 = view_submit(req3)

        rec3 = JHInspectionRecord.objects.get(machine=self.machine_obj, date=date.today(), shift="III")
        self.assert_test(
            rec3.status == JHInspectionRecord.Status.CORRECTED and rec3.corrected_items == 1 and rec3.not_ok_items == 0,
            "TC-JH-06E",
            "Shift III status correctly calculated as CORRECTED (1 Corrected, 0 unresolved)"
        )

        # 6.4 Idempotency / Update or Create Check
        payload_re_edit = {
            "machine_id": self.machine_obj.id,
            "date": today_str,
            "shift": "I",
            "overall_remarks": "Updated remarks on Shift I",
            "results": [
                {"item_id": item1.id, "status": "OK"},
                {"item_id": item2.id, "status": "OK"},
                {"item_id": item3.id, "status": "OK"},
            ]
        }
        req_re = self.factory.post('/api/inspections/jh/submit/', payload_re_edit, format='json')
        force_authenticate(req_re, user=self.operator_user)
        res_re = view_submit(req_re)

        shift_i_count = JHInspectionRecord.objects.filter(machine=self.machine_obj, date=date.today(), shift="I").count()
        self.assert_test(
            shift_i_count == 1,
            "TC-JH-06F",
            "Duplicate submission updates existing record without duplicate key error",
            f"Records count: {shift_i_count}"
        )
        rec1.refresh_from_db()
        self.assert_test(
            rec1.overall_remarks == "Updated remarks on Shift I",
            "TC-JH-06G",
            "Updated remarks properly persisted upon resubmission"
        )

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 7: REPORTS & DETAIL VIEWS
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase7_reports_and_history(self):
        print("\n--- PHASE 7: Inspection Reports & History Filtering ---")

        view_reports = JHInspectionReportsView.as_view()

        # 7.1 Filter by machine
        req_filter = self.factory.get(f'/api/inspections/jh/reports/?machine={self.machine_obj.id}')
        force_authenticate(req_filter, user=self.admin_user)
        res_filter = view_reports(req_filter)

        filter_results = res_filter.data.get("results", [])
        self.assert_test(
            res_filter.status_code == 200 and len(filter_results) >= 3,
            "TC-JH-07A",
            "GET /api/inspections/jh/reports/ returns inspection history for machine",
            f"Returned count: {len(filter_results)}"
        )

        # 7.2 Filter by status
        req_status = self.factory.get(f'/api/inspections/jh/reports/?machine={self.machine_obj.id}&status=HAS_ISSUES')
        force_authenticate(req_status, user=self.admin_user)
        res_status = view_reports(req_status)
        status_results = res_status.data.get("results", [])

        self.assert_test(
            len(status_results) >= 1 and all(r["status"] == "HAS_ISSUES" for r in status_results),
            "TC-JH-07B",
            "Reports endpoint filters accurately by status (HAS_ISSUES)",
            f"Matching records: {len(status_results)}"
        )

        # 7.3 Detail View
        rec = JHInspectionRecord.objects.filter(machine=self.machine_obj).first()
        view_detail = JHInspectionDetailView.as_view()
        req_detail = self.factory.get(f'/api/inspections/jh/reports/{rec.id}/')
        force_authenticate(req_detail, user=self.admin_user)
        res_detail = view_detail(req_detail, pk=str(rec.id))

        self.assert_test(
            res_detail.status_code == 200,
            "TC-JH-07C",
            "GET /api/inspections/jh/reports/<pk>/ returns single record detail"
        )
        self.assert_test(
            "item_results" in res_detail.data and len(res_detail.data["item_results"]) == 3,
            "TC-JH-07D",
            "Record detail contains full nested item_results evaluations"
        )

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 8: MONTHLY 31-DAY MONITORING MATRIX API
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase8_matrix_grid_api(self):
        print("\n--- PHASE 8: Monthly 31-Day Monitoring Matrix Grid ---")

        today = date.today()
        month_str = f"{today.year:04d}-{today.month:02d}"
        view_matrix = JHInspectionMatrixView.as_view()

        # 8.1 8-Hour Schedule (3 Shifts per day)
        req_matrix = self.factory.get(
            f'/api/inspections/jh/matrix/?machine={self.machine_obj.id}&month={month_str}&shift_hours=8'
        )
        force_authenticate(req_matrix, user=self.admin_user)
        res_matrix = view_matrix(req_matrix)

        self.assert_test(
            res_matrix.status_code == 200,
            "TC-JH-08A",
            "GET /api/inspections/jh/matrix/ returns HTTP 200"
        )
        data = res_matrix.data
        self.assert_test(
            data.get("shifts") == ["I", "II", "III"],
            "TC-JH-08B",
            "Matrix returns 3 shifts (['I', 'II', 'III']) for 8-hour schedule"
        )

        # Verify today's evaluations exist in matrix
        matrix_grid = data.get("matrix", {})
        col_shift_1 = f"{today.day}_I"
        col_shift_2 = f"{today.day}_II"
        col_shift_3 = f"{today.day}_III"

        val_1 = matrix_grid.get("TEST-1.1", {}).get(col_shift_1, {})
        val_status_1 = val_1.get("status") if isinstance(val_1, dict) else val_1
        self.assert_test(
            val_status_1 == "OK",
            "TC-JH-08C",
            f"Matrix grid maps Shift I for TEST-1.1 as OK ({col_shift_1})"
        )

        val_2 = matrix_grid.get("TEST-1.2", {}).get(col_shift_2, {})
        val_status_2 = val_2.get("status") if isinstance(val_2, dict) else val_2
        self.assert_test(
            val_status_2 == "NOT_OK",
            "TC-JH-08D",
            f"Matrix grid maps Shift II for TEST-1.2 as NOT_OK ({col_shift_2})"
        )

        val_3 = matrix_grid.get("TEST-1.2", {}).get(col_shift_3, {})
        val_status_3 = val_3.get("status") if isinstance(val_3, dict) else val_3
        self.assert_test(
            val_status_3 == "CORRECTED",
            "TC-JH-08E",
            f"Matrix grid maps Shift III for TEST-1.2 as CORRECTED ({col_shift_3})"
        )

        # 8.2 12-Hour Schedule (2 Shifts per day)
        req_12h = self.factory.get(
            f'/api/inspections/jh/matrix/?machine={self.machine_obj.id}&month={month_str}&shift_hours=12'
        )
        force_authenticate(req_12h, user=self.admin_user)
        res_12h = view_matrix(req_12h)
        self.assert_test(
            res_12h.data.get("shifts") == ["I", "II"],
            "TC-JH-08F",
            "Matrix returns 2 shifts (['I', 'II']) when shift_hours=12"
        )

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 9: FINAL EXCEL EXPORT (FORM QF/MF-08)
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase9_final_excel_export(self):
        print("\n--- PHASE 9: Final Official Form QF/MF-08 Excel Export ---")

        today = date.today()

        # 9.1 Generator direct check
        excel_buf = generate_jh_matrix_xlsx(
            machine=self.machine_obj,
            year=today.year,
            month=today.month,
            shift_hours=8
        )
        self.assert_test(
            len(excel_buf.getvalue()) > 5000,
            "TC-JH-09A",
            "generate_jh_matrix_xlsx produces non-empty workbook buffer (>5 KB)",
            f"Size: {len(excel_buf.getvalue())} bytes"
        )

        # Inspect generated workbook with openpyxl
        wb = openpyxl.load_workbook(filename=excel_buf)
        ws = wb.active
        self.assert_test(
            ws.title == f"JH Matrix {today.year}-{today.month:02d}",
            "TC-JH-09B",
            f"Workbook sheet title is formatted as 'JH Matrix {today.year}-{today.month:02d}'"
        )

        # Check official headers and badges
        row1_text = " ".join([str(ws.cell(row=1, column=c).value or '') for c in range(1, 20)])
        self.assert_test(
            "JISHU-HOZEN" in row1_text and "BREAK DOWN" in row1_text and "ACCIDENT" in row1_text,
            "TC-JH-09C",
            "Workbook Row 1 contains JISHU-HOZEN title and warning badges"
        )

        row2_text = " ".join([str(ws.cell(row=2, column=c).value or '') for c in range(1, 20)])
        self.assert_test(
            self.machine_obj.machine_code in row2_text and "JH STATUS :-  STEP 3" in row2_text,
            "TC-JH-09D",
            f"Workbook Row 2 contains Machine Code ({self.machine_obj.machine_code}) and JH STEP-3"
        )

        # Check evaluation marks in cells (checkmark for OK, cross for NOT OK, circle-cross for CORRECTED)
        all_cell_values = []
        for r in range(5, ws.max_row + 1):
            for c in range(16, ws.max_column + 1):
                val = str(ws.cell(row=r, column=c).value or '')
                if val:
                    all_cell_values.append(val)

        has_checkmark = any("\u2713" in v for v in all_cell_values)
        has_cross = any("\u2715" in v for v in all_cell_values)
        has_corrected = any("\u2297" in v for v in all_cell_values)

        self.assert_test(
            has_checkmark,
            "TC-JH-09E",
            "Workbook cell evaluations contain official green checkmark [OK]"
        )
        self.assert_test(
            has_cross,
            "TC-JH-09F",
            "Workbook cell evaluations contain official red cross [NOT_OK]"
        )
        self.assert_test(
            has_corrected,
            "TC-JH-09G",
            "Workbook cell evaluations contain official blue circle-cross [CORRECTED]"
        )

        # 9.2 HTTP Export Endpoint
        view_export = JHInspectionMatrixExportExcelView.as_view()
        req_export = self.factory.get(
            f'/api/inspections/jh/matrix/export_excel/?machine={self.machine_obj.id}&month={today.month}&year={today.year}'
        )
        force_authenticate(req_export, user=self.admin_user)
        res_export = view_export(req_export)

        self.assert_test(
            res_export.status_code == 200,
            "TC-JH-09H",
            "GET /api/inspections/jh/matrix/export_excel/ returns HTTP 200"
        )
        self.assert_test(
            f"JH_Matrix_{self.machine_obj.machine_code}" in res_export.get('Content-Disposition', ''),
            "TC-JH-09I",
            "Export response filename contains machine code and month"
        )

    # ──────────────────────────────────────────────────────────────────────────
    # PHASE 10: MULTI-MACHINE CHECKLIST ISOLATION & PROVISIONING
    # ──────────────────────────────────────────────────────────────────────────
    def test_phase10_multimachine_isolation(self):
        print("\n--- PHASE 10: Multi-Machine Isolation & Status Provisioning ---")

        # 10.1 Create new machine 2 (JH-MCH-02) and machine 3 (JH-MCH-03)
        mch2, _ = Machine.objects.get_or_create(
            machine_code="JH-MCH-02",
            plant=self.plant_obj,
            defaults={"name": "VMC Milling Center 02", "status": "active"}
        )
        mch3, _ = Machine.objects.get_or_create(
            machine_code="JH-MCH-03",
            plant=self.plant_obj,
            defaults={"name": "Grinding Machine 03", "status": "active"}
        )

        # 10.2 Verify initial status: mch2 has NO custom checklist (inherits default template)
        view_status = JHChecklistStatusView.as_view()
        req_stat = self.factory.get(f'/api/inspections/jh/checklist/status/?machine={mch2.id}')
        force_authenticate(req_stat, user=self.admin_user)
        res_stat = view_status(req_stat)

        self.assert_test(
            res_stat.status_code == 200,
            "TC-JH-10A",
            "GET /api/inspections/jh/checklist/status/ returns HTTP 200"
        )
        self.assert_test(
            res_stat.data.get('is_default') is True and res_stat.data.get('has_custom_checklist') is False,
            "TC-JH-10B",
            "Newly registered machine initially inherits factory default checklist",
            f"Data: {res_stat.data}"
        )

        # 10.3 Bulk Save a custom 2-item checklist specifically for JH-MCH-02
        view_bulk = JHChecklistBulkSaveView.as_view()
        mch2_items = [
            {
                "sub_no": "MCH2-1.1",
                "assembly": "VMC Spindle",
                "check_point": "स्पिंडल कूलेंट प्रेशर चेक करें",
                "standard": "प्रेशर 4.0 बार होना चाहिए",
                "tool_type": "VISUAL",
                "rank": "A",
                "frequency": "D",
                "sort_order": 1
            },
            {
                "sub_no": "MCH2-1.2",
                "assembly": "VMC Tool Changer",
                "check_point": "टूल पॉकेट ग्रीसिंग और क्लैंपिंग चेक करें",
                "standard": "साफ और चिकनाई युक्त",
                "tool_type": "TOUCH",
                "rank": "B",
                "frequency": "D",
                "sort_order": 2
            }
        ]
        payload_mch2 = {
            "machine_id": mch2.id,
            "replace_all": True,
            "filename": "vmc_02_checklist.xlsx",
            "items": mch2_items
        }
        req_save_m2 = self.factory.post('/api/inspections/jh/checklist/bulk_save/', payload_mch2, format='json')
        force_authenticate(req_save_m2, user=self.admin_user)
        res_save_m2 = view_bulk(req_save_m2)

        self.assert_test(
            res_save_m2.status_code == 200 and res_save_m2.data.get('version_number') == 1,
            "TC-JH-10C",
            "Bulk save for JH-MCH-02 creates Version v1 scoped to JH-MCH-02"
        )

        # 10.4 Verify JH-MCH-02 status is now custom active
        req_stat2 = self.factory.get(f'/api/inspections/jh/checklist/status/?machine={mch2.id}')
        force_authenticate(req_stat2, user=self.admin_user)
        res_stat2 = view_status(req_stat2)

        self.assert_test(
            res_stat2.data.get('has_custom_checklist') is True and res_stat2.data.get('total_items') == 2,
            "TC-JH-10D",
            "JH-MCH-02 status reflects 2 custom active checkpoints",
            f"Items: {res_stat2.data.get('total_items')}"
        )

        # 10.5 Verify Default template checkpoints are completely unaffected by JH-MCH-02 upload
        default_items_count = JHChecklistItem.objects.filter(machine__isnull=True, is_active=True, sub_no__startswith="TEST-").count()
        self.assert_test(
            default_items_count == 3,
            "TC-JH-10E",
            "Multi-Machine Isolation: Default template retained all 3 active checkpoints without modification",
            f"Default count: {default_items_count}"
        )

        # 10.5B Save custom checklist for JH-MCH-01 and verify JH-MCH-02 is unaffected
        payload_mch1 = {
            "machine_id": self.machine_obj.id,
            "replace_all": True,
            "filename": "cnc_01_custom.xlsx",
            "items": [
                {
                    "sub_no": "MCH1-1.1",
                    "assembly": "CNC Spindle",
                    "check_point": "चेक चक क्लैंपिंग प्रेशर",
                    "standard": "25 बार",
                    "tool_type": "VISUAL",
                    "rank": "A",
                    "frequency": "D",
                    "sort_order": 1
                }
            ]
        }
        req_save_m1 = self.factory.post('/api/inspections/jh/checklist/bulk_save/', payload_mch1, format='json')
        force_authenticate(req_save_m1, user=self.admin_user)
        res_save_m1 = view_bulk(req_save_m1)
        self.assert_test(
            res_save_m1.status_code == 200,
            "TC-JH-10E2",
            "Custom checklist saved for JH-MCH-01"
        )
        # Verify JH-MCH-02 still has its 2 items active
        mch2_active_count = JHChecklistItem.objects.filter(machine=mch2, is_active=True).count()
        self.assert_test(
            mch2_active_count == 2,
            "TC-JH-10E3",
            "Multi-Machine Isolation: JH-MCH-02 retained its 2 checkpoints when JH-MCH-01 was updated",
            f"MCH-02 count: {mch2_active_count}"
        )

        # 10.6 Retrieve items for JH-MCH-02 (Mobile endpoint)
        view_items = JHChecklistItemsView.as_view()
        req_items_m2 = self.factory.get(f'/api/inspections/jh/items/?machine={mch2.id}')
        force_authenticate(req_items_m2, user=self.operator_user)
        res_items_m2 = view_items(req_items_m2)

        self.assert_test(
            res_items_m2.data.get('count') == 2 and res_items_m2.data.get('is_custom') is True,
            "TC-JH-10F",
            "GET /api/inspections/jh/items/?machine=MCH-02 returns exactly its 2 custom checkpoints"
        )

        # 10.7 Retrieve items for JH-MCH-03 (Fallback to default template)
        req_items_m3 = self.factory.get(f'/api/inspections/jh/items/?machine={mch3.id}')
        force_authenticate(req_items_m3, user=self.operator_user)
        res_items_m3 = view_items(req_items_m3)

        self.assert_test(
            res_items_m3.status_code == 200 and res_items_m3.data.get('is_custom') is False,
            "TC-JH-10G",
            "Machine with no custom checklist (JH-MCH-03) safely inherits default checkpoints (is_custom=False)"
        )

        # 10.8 Template download customized with machine code
        view_tmpl = JHChecklistTemplateDownloadView.as_view()
        req_tmpl = self.factory.get(f'/api/inspections/jh/checklist/template/?machine={mch2.machine_code}')
        force_authenticate(req_tmpl, user=self.admin_user)
        res_tmpl = view_tmpl(req_tmpl)

        self.assert_test(
            "JH_Checklist_Template_JH-MCH-02.xlsx" in res_tmpl.get('Content-Disposition', ''),
            "TC-JH-10H",
            "Template download customized with target machine code in filename"
        )

        # 10.9 Matrix endpoint for JH-MCH-02 reflects its 2 custom checkpoints
        view_mat = JHInspectionMatrixView.as_view()
        req_mat_m2 = self.factory.get(f'/api/inspections/jh/matrix/?machine={mch2.id}&month={date.today().strftime("%Y-%m")}')
        force_authenticate(req_mat_m2, user=self.admin_user)
        res_mat_m2 = view_mat(req_mat_m2)

        mat_items_m2 = res_mat_m2.data.get('items', [])
        self.assert_test(
            len(mat_items_m2) == 2 and mat_items_m2[0]['sub_no'] == 'MCH2-1.1',
            "TC-JH-10I",
            "31-Day Matrix for JH-MCH-02 returns exactly its 2 custom checkpoint rows",
            f"Matrix items: {len(mat_items_m2)}"
        )

        # Clean up temporary test machines
        JHChecklistItem.objects.filter(machine__in=[mch2, mch3]).delete()
        JHChecklistVersion.objects.filter(machine__in=[mch2, mch3]).delete()
        mch2.delete()
        mch3.delete()


if __name__ == '__main__':
    runner = JHTestRunner()
    success = runner.run_all()
    sys.exit(0 if success else 1)
