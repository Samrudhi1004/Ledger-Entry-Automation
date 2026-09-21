import calendar
from datetime import date
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from django.db.models import Prefetch

from apps.inspections.models import DailyProductionReport, DowntimeReport
from apps.machines.models import Machine
from apps.parts.models import InspectionTemplate

def generate_monthly_oee_excel(machine_code: str, year: int, month: int) -> BytesIO:
    """
    Generates a monthly OEE report Excel file for the given machine, year, and month.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = machine_code

    machine = Machine.objects.filter(machine_code=machine_code).first()
    company_name = "AUTO ASSEMBLY CENTRE"
    if machine and machine.plant and hasattr(machine.plant, 'factory') and machine.plant.factory:
        company_name = machine.plant.factory.name
    elif machine and machine.plant:
        company_name = machine.plant.name

    month_name = calendar.month_abbr[month].upper()
    
    # Define Styles
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    def set_cell(row, col, value, font=None, alignment=center_align, border=thin_border, fill=None):
        cell = ws.cell(row=row, column=col, value=value)
        if font: cell.font = font
        if alignment: cell.alignment = alignment
        if border: cell.border = border
        if fill: cell.fill = fill
        return cell

    # Row 1: Company Name & Title
    ws.merge_cells('A1:D1')
    set_cell(1, 1, company_name, font=Font(bold=True, size=14), alignment=Alignment(horizontal='left'))
    ws.merge_cells('E1:M1')
    set_cell(1, 5, "O.E.E. SHEET", font=Font(bold=True, size=16))
    ws.merge_cells('N1:T1')
    set_cell(1, 14, "Form No.: SGF/AAC/PRD/F/06", alignment=Alignment(horizontal='right'))

    # Row 2: Month & Machine
    ws.merge_cells('A2:H2')
    set_cell(2, 1, f"MONTH - {month_name} {year}", font=bold_font, alignment=Alignment(horizontal='left'))
    ws.merge_cells('I2:T2')
    set_cell(2, 9, f"MACHINE - {machine_code}", font=bold_font, alignment=Alignment(horizontal='left'))

    # Define Header Matrix (Rows 3, 4, 5, 6)
    headers = [
        # Col 1: A
        ("SR.NO.", "", "", ""),
        # Col 2: B
        ("DATE", "", "", ""),
        # Col 3: C
        ("AVAILABLE TIME", "A", "MINUTES", ""),
        # Col 4: D
        ("PLANNED DOWNTIME", "B", "MINUTES", ""),
        # Col 5: E
        ("NET AVAILABLE TIME", "C", "MINUTES", "C=A-B"),
        # Col 6: F
        ("DOWN TIME LOSSES", "D", "MINUTES", ""),
        # Col 7: G
        ("DOWN TIME RES.", "ST", "", ""),
        # Col 8: H
        ("", "NL", "", ""),
        # Col 9: I
        ("", "NO", "", ""),
        # Col 10: J
        ("", "MM", "", ""),
        # Col 11: K
        ("", "OW", "", ""),
        # Col 12: L
        ("", "PF", "", ""),
        # Col 13: M
        ("OPERATING TIME", "E", "MINUTES", "E=C-D"),
        # Col 14: N
        ("AVAILABILITY", "F", "%", "F=E/C"),
        # Col 15: O
        ("TOTAL QUANTITY PRODUCED", "G", "BATCH (nos)", ""),
        # Col 16: P
        ("THEORETICAL CYCLE TIME", "H", "MINUTES", ""),
        # Col 17: Q
        ("PERFORMANCE EFFICIENCY", "I", "%", "I=(G*H/E)"),
        # Col 18: R
        ("REJECTION", "J", "NOS.", ""),
        # Col 19: S
        ("RATE OF QUALITY PRODUCT", "K", "%", "K=(G-J)/G"),
        # Col 20: T
        ("O.E.E", "L", "%", "L=F*I*K*100"),
    ]

    # Write Headers (Rows 3 to 6)
    for col_idx, (r3, r4, r5, r6) in enumerate(headers, start=1):
        set_cell(3, col_idx, r3, font=bold_font, fill=header_fill)
        set_cell(4, col_idx, r4, font=bold_font, fill=header_fill)
        set_cell(5, col_idx, r5, font=bold_font, fill=header_fill)
        set_cell(6, col_idx, r6, font=bold_font, fill=header_fill)

    # Set Column Widths
    column_widths = {
        'A': 8, 'B': 12, 'C': 12, 'D': 12, 'E': 12, 'F': 12,
        'G': 6, 'H': 6, 'I': 6, 'J': 6, 'K': 6, 'L': 6,
        'M': 12, 'N': 12, 'O': 12, 'P': 12, 'Q': 12, 'R': 10, 'S': 12, 'T': 12
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Fetch Data
    reports = DailyProductionReport.objects.filter(
        machine__machine_code=machine_code,
        date__year=year,
        date__month=month
    ).select_related('downtime_report', 'part').order_by('date', 'shift')

    current_row = 7
    last_date = None

    for report in reports:
        dt = getattr(report, 'downtime_report', None)
        
        # Calculate available and planned downtime
        available_time = 480
        planned_downtime = 60
        if machine and machine.plant:
            if hasattr(machine.plant, 'factory') and machine.plant.factory and machine.plant.factory.shift_hours:
                fac = machine.plant.factory
                available_time = fac.shift_hours * 60
                planned_downtime = fac.lunch_break_minutes + fac.tea_break_minutes
            elif machine.plant.shift_duration_hours:
                available_time = machine.plant.shift_duration_hours * 60
                planned_downtime = machine.plant.total_break_mins or 0

        # Cycle Time
        cycle_time = 0.0
        if report.part:
            template = InspectionTemplate.objects.filter(part=report.part, name=report.operation).first()
            if not template:
                template = InspectionTemplate.objects.filter(part=report.part).first()
            if template and getattr(template, 'cycle_time_mins', 0) > 0:
                cycle_time = template.cycle_time_mins

        # A: SR NO / Shift
        set_cell(current_row, 1, report.shift)
        
        # B: DATE
        date_val = report.date.strftime("%Y-%m-%d") if report.date != last_date else ""
        set_cell(current_row, 2, date_val)
        last_date = report.date

        # C: Available Time (A)
        set_cell(current_row, 3, available_time)
        
        # D: Planned Downtime (B)
        set_cell(current_row, 4, planned_downtime)
        
        # E: Net Available (C) -> Formula: C - D  (Which is A - B in excel headers)
        ws.cell(row=current_row, column=5).value = f"=C{current_row}-D{current_row}"
        ws.cell(row=current_row, column=5).alignment = center_align
        ws.cell(row=current_row, column=5).border = thin_border
        
        # F: Down Time Losses (D) -> Sum of ST to PF (G to L)
        ws.cell(row=current_row, column=6).value = f"=SUM(G{current_row}:L{current_row})"
        ws.cell(row=current_row, column=6).alignment = center_align
        ws.cell(row=current_row, column=6).border = thin_border
        
        # Downtime reasons (G to L)
        set_cell(current_row, 7, dt.setting if dt and dt.setting else "")        # ST
        set_cell(current_row, 8, dt.no_load if dt and dt.no_load else "")        # NL
        set_cell(current_row, 9, dt.no_operator if dt and dt.no_operator else "")# NO
        set_cell(current_row, 10, dt.um if dt and dt.um else "")                 # MM (um)
        set_cell(current_row, 11, dt.inspection_wait if dt and dt.inspection_wait else "") # OW
        set_cell(current_row, 12, dt.power_off if dt and dt.power_off else "")   # PF
        
        # M: Operating Time (E) -> Formula: C - D (E - F in excel cols)
        ws.cell(row=current_row, column=13).value = f"=E{current_row}-F{current_row}"
        ws.cell(row=current_row, column=13).alignment = center_align
        ws.cell(row=current_row, column=13).border = thin_border
        
        # N: Availability (F) -> Formula: E / C (M / E in excel cols)
        ws.cell(row=current_row, column=14).value = f"=IF(E{current_row}>0, M{current_row}/E{current_row}, 0)"
        ws.cell(row=current_row, column=14).alignment = center_align
        ws.cell(row=current_row, column=14).border = thin_border
        ws.cell(row=current_row, column=14).number_format = '0.00%'
        
        # O: Total Qty (G)
        set_cell(current_row, 15, report.jobs_completed or 0)
        
        # P: Theo Cycle Time (H)
        set_cell(current_row, 16, round(cycle_time, 5))
        
        # Q: Performance (I) -> Formula: (G*H)/E  ((O*P)/M in excel cols)
        ws.cell(row=current_row, column=17).value = f"=IF(M{current_row}>0, (O{current_row}*P{current_row})/M{current_row}, 0)"
        ws.cell(row=current_row, column=17).alignment = center_align
        ws.cell(row=current_row, column=17).border = thin_border
        ws.cell(row=current_row, column=17).number_format = '0.00%'
        
        # R: Rejection (J)
        set_cell(current_row, 18, report.incorrect_jobs or 0)
        
        # S: Rate of Quality (K) -> Formula: (G-J)/G ((O-R)/O in excel cols)
        ws.cell(row=current_row, column=19).value = f"=IF(O{current_row}>0, (O{current_row}-R{current_row})/O{current_row}, 0)"
        ws.cell(row=current_row, column=19).alignment = center_align
        ws.cell(row=current_row, column=19).border = thin_border
        ws.cell(row=current_row, column=19).number_format = '0.00%'
        
        # T: O.E.E (L) -> Formula: F*I*K*100 (N*Q*S in percentage format)
        # Note: If we use percentage format, we don't need * 100
        ws.cell(row=current_row, column=20).value = f"=N{current_row}*Q{current_row}*S{current_row}"
        ws.cell(row=current_row, column=20).alignment = center_align
        ws.cell(row=current_row, column=20).border = thin_border
        ws.cell(row=current_row, column=20).number_format = '0.00%'

        current_row += 1

    # Freeze panes below headers
    ws.freeze_panes = 'A7'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
